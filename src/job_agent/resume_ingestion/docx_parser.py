from pathlib import Path

from docx import Document

from job_agent.resume_ingestion.parsers import ParsedPage, ParsedResume, hash_file


def parse_docx(path: Path) -> ParsedResume:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedResume(
        filename=path.name,
        file_hash=hash_file(path),
        extraction_method="docx",
        pages=[ParsedPage(page_number=1, text="\n".join(blocks))],
    )
